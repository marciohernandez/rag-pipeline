"""Testes da normalização de formatos Office legados (SPEC-0023)."""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

FIXTURES = Path(__file__).parent / "fixtures" / "spec0023"


def _normalizar_com_soffice_simulado(
    tmp_path,
    monkeypatch,
    *,
    extensao_origem: str,
    extensao_destino: str,
) -> None:
    from src.pipeline.conversion.legacy_office import normalizar_formato_legado

    origem = tmp_path / f"documento.{extensao_origem}"
    origem.write_bytes(b"fixture-legada")
    comandos: list[list[str]] = []

    def executar(comando, **_kwargs):
        argumentos = [str(item) for item in comando]
        comandos.append(argumentos)
        pasta_saida = Path(argumentos[argumentos.index("--outdir") + 1])
        convertido = pasta_saida / f"{origem.stem}.{extensao_destino}"
        if extensao_destino == "docx":
            shutil.copyfile(FIXTURES / "estrutura.docx", convertido)
        elif extensao_destino == "xlsx":
            shutil.copyfile(FIXTURES / "duas-planilhas.xlsx", convertido)
        else:
            from pptx import Presentation

            apresentacao = Presentation()
            slide = apresentacao.slides.add_slide(apresentacao.slide_layouts[0])
            slide.shapes.title.text = "Apresentação legada"
            apresentacao.save(convertido)
        return subprocess.CompletedProcess(argumentos, 0, stdout="", stderr="")

    monkeypatch.setattr(subprocess, "run", executar)

    convertido = normalizar_formato_legado(origem, extensao_origem)

    assert convertido.suffix == f".{extensao_destino}"
    assert convertido.stat().st_size > 0
    assert comandos
    assert comandos[0][:3] == ["soffice", "--headless", "--convert-to"]
    assert comandos[0][3] == extensao_destino
    assert "--outdir" in comandos[0]
    assert comandos[0][-1] == str(origem)

    if extensao_destino == "docx":
        from src.pipeline.conversion.converters.docx import DocxConverter

        markdown = DocxConverter().converter(convertido).markdown
        assert "Relatório de atendimento" in markdown
    elif extensao_destino == "pptx":
        from src.pipeline.conversion.converters.pptx import PptxConverter

        markdown = PptxConverter().converter(convertido).markdown
        assert "Apresentação legada" in markdown
    else:
        from src.pipeline.conversion.converters.xlsx import XlsxConverter

        markdown = XlsxConverter().converter(convertido).markdown
        assert "Resumo" in markdown


def test_spec_0023_legacy_doc_e_normalizado_para_docx(tmp_path, monkeypatch) -> None:
    """SPECSFY: SPEC-0023 US-001 FR-003"""
    _normalizar_com_soffice_simulado(
        tmp_path,
        monkeypatch,
        extensao_origem="doc",
        extensao_destino="docx",
    )


def test_spec_0023_legacy_ppt_e_normalizado_para_pptx(tmp_path, monkeypatch) -> None:
    """SPECSFY: SPEC-0023 US-001 FR-003"""
    _normalizar_com_soffice_simulado(
        tmp_path,
        monkeypatch,
        extensao_origem="ppt",
        extensao_destino="pptx",
    )


def test_spec_0023_legacy_xls_e_normalizado_para_xlsx(tmp_path, monkeypatch) -> None:
    """SPECSFY: SPEC-0023 US-001 FR-003"""
    _normalizar_com_soffice_simulado(
        tmp_path,
        monkeypatch,
        extensao_origem="xls",
        extensao_destino="xlsx",
    )
