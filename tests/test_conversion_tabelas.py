"""Tabelas simples, complexas e contínuas entre páginas (SPEC-0040).

Contrato assumido: ``src.pipeline.conversion.tabelas`` expõe ``Texto(conteudo)``,
``Tabela(linhas, cabecalho_niveis=1, mesclada=False)`` e ``normalizar_tabelas(paginas) ->
ResultadoTabelas(markdown, avisos)``, em que ``paginas`` é uma lista (uma por página) de listas de
blocos. Tabela simples vira tabela Markdown; complexa (cabeçalho de vários níveis ou células
mescladas) vira texto corrido, uma frase por linha, com o aviso ``tabela_complexa_virou_texto``;
tabelas que terminam e começam páginas consecutivas com as mesmas colunas são juntadas (sem repetir
o cabeçalho); com colunas diferentes ficam separadas com o aviso ``tabela_nao_juntada``.
"""

from __future__ import annotations

import time

from tests._fase1_conversao import exigir

CABECALHO = ["Item", "Prazo", "Valor"]


def _tabela(modulo, *linhas, **extra):
    return modulo.Tabela([list(CABECALHO), *map(list, linhas)], **extra)


def _codigos(resultado) -> list[str]:
    return [a.codigo for a in resultado.avisos]


def _linhas_de_tabela(markdown: str) -> list[str]:
    return [linha for linha in markdown.splitlines() if linha.startswith("|")]


def test_spec_0040_ac_012_tabela_simples_vira_tabela_markdown() -> None:
    """SPECSFY: SPEC-0040 US-004 FR-005 AC-012"""
    modulo = exigir("src.pipeline.conversion.tabelas")
    tabela = _tabela(modulo, ["Plano básico", "7 dias", "R$ 50"], ["Plano pro", "7 dias", "R$ 90"])

    resultado = modulo.normalizar_tabelas([[modulo.Texto("Preços"), tabela]])

    linhas = _linhas_de_tabela(resultado.markdown)
    assert linhas[0] == "| Item | Prazo | Valor |"
    assert linhas[1].replace(" ", "") == "|---|---|---|"
    assert "| Plano pro | 7 dias | R$ 90 |" in linhas
    assert "Preços" in resultado.markdown
    assert resultado.avisos == []


def test_spec_0040_ac_013_cabecalho_de_dois_niveis_vira_texto_e_marca_revisar() -> None:
    """SPECSFY: SPEC-0040 US-004 FR-005 AC-013"""
    modulo = exigir("src.pipeline.conversion.tabelas")
    tabela = modulo.Tabela(
        [["Plano", "Preço", ""], ["", "Mensal", "Anual"], ["Básico", "R$ 50", "R$ 500"]],
        cabecalho_niveis=2,
    )

    resultado = modulo.normalizar_tabelas([[tabela]])

    assert _linhas_de_tabela(resultado.markdown) == []
    assert "Básico" in resultado.markdown and "R$ 500" in resultado.markdown
    assert "tabela_complexa_virou_texto" in _codigos(resultado)
    assert any("tabela complexa virou texto" in a.mensagem for a in resultado.avisos)


def test_spec_0040_ac_013_texto_corrido_tem_uma_frase_por_linha() -> None:
    """SPECSFY: SPEC-0040 US-004 FR-005 AC-013"""
    modulo = exigir("src.pipeline.conversion.tabelas")
    tabela = _tabela(modulo, ["Plano básico", "7 dias", "R$ 50"], ["Plano pro", "7 dias", "R$ 90"])
    complexa = modulo.Tabela(tabela.linhas, cabecalho_niveis=2)

    resultado = modulo.normalizar_tabelas([[complexa]])

    frases = [linha for linha in resultado.markdown.splitlines() if linha.strip()]
    assert len(frases) >= 2
    assert any("Plano pro" in f and "R$ 90" in f for f in frases)


def test_spec_0040_ac_014_celulas_mescladas_contam_como_complexa(tmp_path) -> None:
    """SPECSFY: SPEC-0040 US-004 FR-005 AC-014"""
    from docx import Document

    from src.pipeline.conversion.converters.docx import DocxConverter

    exigir("src.pipeline.conversion.tabelas")
    documento = Document()
    tabela = documento.add_table(rows=3, cols=3)
    tabela.cell(0, 0).merge(tabela.cell(0, 1))
    tabela.cell(0, 0).text = "Plano e prazo"
    tabela.cell(0, 2).text = "Valor"
    for linha, valores in enumerate((("Básico", "7 dias", "R$ 50"), ("Pro", "7 dias", "R$ 90")), 1):
        for coluna, valor in enumerate(valores):
            tabela.cell(linha, coluna).text = valor
    caminho = tmp_path / "planos.docx"
    documento.save(caminho)

    resultado = DocxConverter().converter(caminho)

    assert "tabela_complexa_virou_texto" in [a.codigo for a in resultado.avisos]
    assert "R$ 90" in resultado.markdown
    assert not any(linha.startswith("| ---") for linha in resultado.markdown.splitlines())


def test_spec_0040_ac_015_tabela_que_atravessa_duas_paginas_vira_uma() -> None:
    """SPECSFY: SPEC-0040 US-005 FR-006 AC-015"""
    modulo = exigir("src.pipeline.conversion.tabelas")
    cab = ["A", "B", "C", "D"]
    pagina1 = [modulo.Tabela([cab, ["1", "2", "3", "4"], ["5", "6", "7", "8"]])]
    pagina2 = [modulo.Tabela([["9", "10", "11", "12"], ["13", "14", "15", "16"]])]

    resultado = modulo.normalizar_tabelas([pagina1, pagina2])

    linhas = _linhas_de_tabela(resultado.markdown)
    assert len(linhas) == 2 + 4  # cabeçalho, separador e quatro linhas de dados
    assert "| 13 | 14 | 15 | 16 |" in linhas
    assert resultado.avisos == []


def test_spec_0040_ac_016_colunas_diferentes_nao_juntam_e_marcam_revisar() -> None:
    """SPECSFY: SPEC-0040 US-005 FR-006 AC-016"""
    modulo = exigir("src.pipeline.conversion.tabelas")
    pagina1 = [modulo.Tabela([["A", "B", "C", "D"], ["1", "2", "3", "4"]])]
    pagina2 = [modulo.Tabela([["X", "Y", "Z"], ["7", "8", "9"]])]

    resultado = modulo.normalizar_tabelas([pagina1, pagina2])

    separadores = [l for l in _linhas_de_tabela(resultado.markdown) if set(l) <= set("|- ")]
    assert len(separadores) == 2
    assert "tabela_nao_juntada" in _codigos(resultado)
    assert any("tabela não juntada" in a.mensagem for a in resultado.avisos)


def test_spec_0040_ac_016_tabela_no_meio_da_pagina_nao_e_candidata_a_juncao() -> None:
    """SPECSFY: SPEC-0040 US-005 FR-006 AC-016"""
    modulo = exigir("src.pipeline.conversion.tabelas")
    cab = ["A", "B"]
    pagina1 = [modulo.Tabela([cab, ["1", "2"]]), modulo.Texto("Parágrafo final da página.")]
    pagina2 = [modulo.Tabela([cab, ["3", "4"]])]

    resultado = modulo.normalizar_tabelas([pagina1, pagina2])

    separadores = [l for l in _linhas_de_tabela(resultado.markdown) if set(l) <= set("|- ")]
    assert len(separadores) == 2


def test_spec_0040_ac_017_cabecalho_repetido_na_pagina_seguinte_nao_duplica() -> None:
    """SPECSFY: SPEC-0040 US-005 FR-006 AC-017"""
    modulo = exigir("src.pipeline.conversion.tabelas")
    cab = ["Item", "Prazo", "Valor"]
    pagina1 = [modulo.Tabela([cab, ["a", "1", "10"]])]
    pagina2 = [modulo.Tabela([cab, ["b", "2", "20"]])]

    resultado = modulo.normalizar_tabelas([pagina1, pagina2])

    linhas = _linhas_de_tabela(resultado.markdown)
    assert linhas.count("| Item | Prazo | Valor |") == 1
    assert "| b | 2 | 20 |" in linhas


def test_spec_0040_ac_027_cem_paginas_com_tabelas_continuas_dentro_do_limite() -> None:
    """SPECSFY: SPEC-0040 FR-006 NFR-003 AC-027"""
    modulo = exigir("src.pipeline.conversion.tabelas")
    cab = ["A", "B", "C"]
    paginas = [
        [modulo.Tabela([cab, ["0", "x", "y"]] if p == 0 else [[str(p), "x", "y"]])]
        for p in range(100)
    ]

    inicio = time.perf_counter()
    resultado = modulo.normalizar_tabelas(paginas)
    duracao = time.perf_counter() - inicio

    assert duracao < 5
    assert len(_linhas_de_tabela(resultado.markdown)) == 2 + 100
    assert resultado.avisos == []
