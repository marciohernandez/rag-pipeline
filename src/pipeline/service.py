"""Fachada estável dos casos de uso do pipeline.

As operações recebem a configuração não sensível já resolvida. Leitura de
arquivo e preparação do ambiente pertencem aos entrypoints da raiz; resolução
de credenciais continua encapsulada nos adapters de provider.
"""

from __future__ import annotations

import logging as _logging
from collections.abc import Callable as _Callable
from collections.abc import Mapping as _Mapping
from collections.abc import Sequence as _Sequence
from dataclasses import dataclass as _dataclass
from pathlib import Path as _Path
from typing import Any as _Any

from langchain_core.documents import Document as _Document

from src.pipeline import indexing as _indexing
from src.pipeline.assistant import AssistantError as _AssistantError
from src.pipeline.assistant import AssistantReply as _AssistantReply
from src.pipeline.assistant import executar_turno as _executar_turno
from src.pipeline.chunker import chunk_documents as _chunk_documents
from src.pipeline.conversion.config import carregar_config_conversao as _carregar_config_conversao
from src.pipeline.conversion.contracts import ConversionResult as _ConversionResult
from src.pipeline.conversion.contracts import ConversionUserError
from src.pipeline.conversion.factory import ConverterFactory as _ConverterFactory
from src.pipeline.exporter import export_chunks_to_json as _export_chunks_to_json
from src.pipeline.identity import content_hash as _content_hash
from src.pipeline.indexing import build_scope as _build_scope
from src.pipeline.indexing import index_chunks as _index_chunks
from src.pipeline.indexing import index_documents as _index_documents
from src.pipeline.indexing import retrieve as _retrieve
from src.pipeline.loader import load_markdown_documents as _load_markdown_documents
from src.pipeline.manifest import build_manifest as _build_manifest
from src.pipeline.manifest import export_manifest as _export_manifest
from src.pipeline.providers.contracts import SyncCounts as _SyncCounts
from src.pipeline.providers.errors import UnreliableScopeReadError as _UnreliableScopeReadError
from src.pipeline.providers.factory import ProviderFactory as _ProviderFactory
from src.pipeline.quality import evaluate_quality as _evaluate_quality
from src.pipeline.safety import SyncSafetyPolicy as _SyncSafetyPolicy
from src.pipeline.status_gate import filter_by_status as _filter_by_status
from src.pipeline.validation import validate_base as _validate_base

__all__ = [
    "AssistantError",
    "AssistantReply",
    "ConversionUserError",
    "aplicar_regras_limpeza",
    "assistant_reply",
    "avaliar_qualidade_conversao",
    "buscar_url_conversao",
    "convert_document",
    "criar_modelo_chat",
    "deduplicar_documentos_limpeza",
    "delete_scope",
    "describe_collection",
    "index_assistant_knowledge",
    "limpar_secoes_ia",
    "listar_abas_conversao",
    "montar_relatorio_limpeza",
    "preview_chunks",
    "propor_dominio_ia",
    "query_base",
    "reprocessar_conversao_ia",
    "run_pipeline",
    "sync_vector_store",
    "validar_config_conversao",
    "validar_url_conversao",
    "validate_retrieval",
]

AssistantError = _AssistantError
AssistantReply = _AssistantReply


@_dataclass(frozen=True)
class _PipelineRunResult:
    """Resultado estruturado da preparação e sincronização de uma base."""

    documents: tuple[_Document, ...]
    chunks: tuple[_Document, ...]
    manifest: dict[str, _Any]
    manifest_path: str
    excluded_documents: tuple[dict[str, _Any], ...]
    quality_report: _Any
    sync_counts: _SyncCounts | None


def _required_mapping(config: _Mapping[str, _Any], key: str) -> dict[str, _Any]:
    value = config.get(key)
    if value is None:
        raise ValueError(f"Configuração sem {key}.")
    if not isinstance(value, _Mapping):
        raise TypeError(f"Configuração {key} deve ser um objeto.")
    return dict(value)


def _optional_mapping(config: _Mapping[str, _Any], key: str) -> dict[str, _Any]:
    value = config.get(key) or {}
    if not isinstance(value, _Mapping):
        raise TypeError(f"Configuração {key} deve ser um objeto.")
    return dict(value)


def _project_id(config: _Mapping[str, _Any]) -> str:
    project = _required_mapping(config, "projeto")
    value = project.get("nome")
    if not isinstance(value, str) or not value.strip():
        raise ValueError("Configuração sem projeto.nome.")
    return value.strip()


def _runtime_sections(
    config: _Mapping[str, _Any],
) -> tuple[str, dict[str, _Any], dict[str, _Any]]:
    if not isinstance(config, _Mapping):
        raise TypeError("Configuração resolvida deve ser um objeto.")
    return (
        _project_id(config),
        _required_mapping(config, "embedding"),
        _required_mapping(config, "vector_store"),
    )


def _required_path(paths: _Mapping[str, _Any], key: str) -> str:
    value = paths.get(key)
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"Configuração sem caminhos.{key}.")
    return value.strip()


def run_pipeline(
    config: _Mapping[str, _Any],
    *,
    input_dir: str | None = None,
    only_chunks: bool = False,
) -> _PipelineRunResult:
    """Prepara uma base e, quando solicitado, sincroniza seus chunks."""
    project_id, embedding_config, vector_config = _runtime_sections(config)
    paths = _required_mapping(config, "caminhos")
    source_dir = input_dir or _required_path(paths, "input_dir")
    chunks_path = _required_path(paths, "chunks_file")
    manifest_path = str(paths.get("manifest_file") or _Path(chunks_path).with_name("manifest.json"))

    documents = _load_markdown_documents(source_dir)
    markdown_contract = _optional_mapping(config, "contrato_markdown")
    selected, excluded = _filter_by_status(
        documents,
        markdown_contract.get("status_bloqueado", []),
    )
    chunking_config = _optional_mapping(config, "chunking")
    chunks = _chunk_documents(
        selected,
        dict(chunking_config),
        project_id=project_id,
    )

    _export_chunks_to_json(chunks, chunks_path)
    manifest = _build_manifest(
        chunks,
        project_id=project_id,
        collection_name=vector_config.get("collection_name", "rag_docs"),
        embedding_cfg=embedding_config,
        excluded_documents=excluded,
    )
    _export_manifest(manifest, manifest_path)

    quality_config = _optional_mapping(config, "qualidade")
    quality_report = _evaluate_quality(
        chunks,
        dict(chunking_config),
        dict(quality_config),
        manifest=manifest,
        output_dir=_Path(chunks_path).parent,
    )
    counts = None
    if not only_chunks:
        counts = sync_vector_store(config=config, chunks=chunks)

    return _PipelineRunResult(
        documents=tuple(selected),
        chunks=tuple(chunks),
        manifest=manifest,
        manifest_path=manifest_path,
        excluded_documents=tuple(excluded),
        quality_report=quality_report,
        sync_counts=counts,
    )


def preview_chunks(
    config: _Mapping[str, _Any],
    *,
    input_dir: str,
) -> _PipelineRunResult:
    """Calcula os chunks de uma pasta sem gerar embeddings nem tocar o destino.

    Reaproveita a mesma divisão da execução (precedência do frontmatter
    inclusive); o chamador entrega um diretório temporário e o descarta depois.
    """
    return run_pipeline(config, input_dir=input_dir, only_chunks=True)


TEMPO_LIMITE_TESTE_DE_CONEXAO = 15.0
_logger = _logging.getLogger(__name__)


def _testar_qdrant(vector_config: _Mapping[str, _Any]) -> dict[str, _Any]:
    from src.pipeline.providers.vectorstores.qdrant import testar_acesso

    return testar_acesso(str(vector_config.get("url", "")), vector_config.get("api_key"))


def _testar_embeddings(embedding_config: _Mapping[str, _Any]) -> dict[str, _Any]:
    declaradas = embedding_config.get("dimensions")
    try:
        adapter = _indexing.build_embeddings(dict(embedding_config))
        vetor = adapter.embed_documents(["teste de conexão"])[0]
    except Exception as erro:  # noqa: BLE001 - cada SDK de embeddings tem sua hierarquia
        status = getattr(erro, "status_code", None)
        if status in (401, 403):
            mensagem = "O provedor recusou a chave informada."
        elif status == 404:
            mensagem = "O provedor não encontrou o modelo informado."
        elif (
            "connection" in type(erro).__name__.lower() or "timeout" in type(erro).__name__.lower()
        ):
            mensagem = "Não foi possível alcançar o servidor."
        else:
            mensagem = "Não foi possível testar a conexão."
        return {"sucesso": False, "mensagem": mensagem, "dimensoes": None}
    obtidas = len(vetor)
    if declaradas and int(declaradas) != obtidas:
        return {
            "sucesso": False,
            "mensagem": (
                f"O modelo devolveu {obtidas} dimensões; a conexão declara {int(declaradas)}."
            ),
            "dimensoes": obtidas,
        }
    return {"sucesso": True, "mensagem": "Conexão testada com sucesso.", "dimensoes": obtidas}


def testar_conexao(configuracao: _Mapping[str, _Any]) -> dict[str, _Any]:
    """Testa o acesso de uma seção `vector_store` (Qdrant) ou `embedding` (SPEC-0029).

    Não cria coleção nem grava pontos. O resultado traz só mensagens próprias,
    nunca o texto da exceção, para que a chave não vaze; o tempo limite cobre
    clientes que não aceitam timeout curto.
    """
    from concurrent.futures import ThreadPoolExecutor
    from concurrent.futures import TimeoutError as _Esgotado

    qdrant = configuracao.get("provider") == "qdrant"
    funcao = _testar_qdrant if qdrant else _testar_embeddings
    executor = ThreadPoolExecutor(max_workers=1)
    try:
        resultado = executor.submit(funcao, configuracao).result(
            timeout=TEMPO_LIMITE_TESTE_DE_CONEXAO
        )
    except _Esgotado:
        resultado = {"sucesso": False, "mensagem": "Não foi possível alcançar o servidor."}
        if not qdrant:
            resultado["dimensoes"] = None
    finally:
        executor.shutdown(wait=False)
    _logger.info(
        "teste de conexão %s: %s",
        "qdrant" if qdrant else "embeddings",
        "sucesso" if resultado["sucesso"] else "falha",
    )
    return resultado


def describe_collection(vector_config: _Mapping[str, _Any]) -> dict[str, _Any]:
    """Lê, sem criar nada, se a coleção existe e qual dimensão ela declara.

    `describe_scope` devolve mapa vazio tanto para coleção ausente quanto para
    destino inacessível (DEC-015); o chamador trata os dois como inexistente.
    """
    adapter = _indexing.build_vector_store(dict(vector_config))
    descricao = adapter.describe_scope(_build_scope("descricao-de-colecao"))
    existe = bool(descricao.get("exists"))
    return {"exists": existe, "dimensions": descricao.get("dimensions") if existe else None}


def sync_vector_store(
    config: _Mapping[str, _Any],
    chunks: _Sequence[_Document],
) -> _SyncCounts:
    """Sincroniza chunks e devolve as quatro contagens ao chamador."""
    _project, embedding_config, vector_config = _runtime_sections(config)
    if not chunks:
        return _SyncCounts(previstos=0, inseridos=0, mantidos=0, removidos=0)
    paths = _optional_mapping(config, "caminhos")
    return _index_chunks(
        chunks=chunks,
        embedding_cfg=embedding_config,
        vector_cfg=vector_config,
        persist_dir=paths.get("vector_store_dir"),
    )


def sync_documents(
    config: _Mapping[str, _Any],
    chunks: _Sequence[_Document],
    on_document: _Callable[[str, _SyncCounts, list[_Document]], None] | None = None,
) -> _SyncCounts:
    """Sincroniza cada documento no escopo próprio (coleção + documento).

    Somar, manter ou substituir acontece por documento (SPEC-0030); `on_document`
    é chamado depois que o destino confirma cada documento.
    """
    _project, embedding_config, vector_config = _runtime_sections(config)
    paths = _optional_mapping(config, "caminhos")
    return _index_documents(
        chunks=chunks,
        embedding_cfg=embedding_config,
        vector_cfg=vector_config,
        persist_dir=paths.get("vector_store_dir"),
        on_document=on_document,
    )


def list_document_chunks(
    config: _Mapping[str, _Any],
    document_id: str,
    limit: int = 200,
) -> list[dict[str, _Any]]:
    """Lista, somente para leitura, os chunks de um documento da coleção."""
    project_id, _embedding_config, vector_config = _runtime_sections(config)
    adapter = _indexing.build_vector_store(vector_config)
    escopo = _build_scope(project_id, None, document_id)
    try:
        registros = adapter.list_records(escopo, limit)
    except Exception as erro:
        if getattr(erro, "status_code", None) == 404:
            return []
        raise
    return [
        {
            "chunk_id": registro.metadata.get("chunk_id"),
            "secao": registro.metadata.get("secao") or registro.metadata.get("titulo") or "",
            "texto": registro.content,
        }
        for registro in registros
    ]


def drop_collection(config: _Mapping[str, _Any], *, registradas: _Sequence[str]) -> None:
    """Apaga a coleção física do destino; só aceita nomes registrados pelo app."""
    _project, _embedding_config, vector_config = _runtime_sections(config)
    nome = str(vector_config.get("collection_name", ""))
    _SyncSafetyPolicy().assert_collection_drop_allowed(nome, registradas)
    adapter = _indexing.build_vector_store(vector_config)
    adapter.drop()


def delete_scope(
    config: _Mapping[str, _Any],
    *,
    tenant_id: str | None = None,
    document_id: str | None = None,
) -> int:
    """Remove todos os pontos indexados no escopo lógico configurado.

    Usado pela exclusão de bases (`src/app/`): cada base tem seu próprio
    `projeto.nome` (isolamento por base resolvido pelo chamador), então
    esvaziar o escopo aqui não afeta outras bases.
    """
    project_id, _embedding_config, vector_config = _runtime_sections(config)
    adapter = _indexing.build_vector_store(vector_config)
    scope = _build_scope(project_id, tenant_id, document_id)
    try:
        ids = adapter.list_ids(scope)
    except _UnreliableScopeReadError as erro:
        # Coleção ausente (404 do destino) não tem pontos a remover; qualquer
        # outra falha de leitura continua bloqueando a exclusão (SPEC-0028).
        if getattr(erro.__cause__, "status_code", None) != 404:
            raise
        _logging.getLogger(__name__).info(
            "Coleção %s ausente no vector store; nada a remover.",
            vector_config.get("collection_name"),
        )
        return 0
    if ids:
        adapter.delete_ids(scope, ids)
    return len(ids)


def validate_retrieval(
    config: _Mapping[str, _Any],
    *,
    tenant_id: str | None = None,
):
    """Valida uma base indexada sem alterar o destino."""
    _runtime_sections(config)
    _required_mapping(config, "caminhos")
    return _validate_base(dict(config), tenant_id=tenant_id)


def query_base(
    config: _Mapping[str, _Any],
    query: str,
    *,
    top_k: int = 4,
    tenant_id: str | None = None,
):
    """Consulta a base indexada sob o escopo lógico configurado."""
    project_id, embedding_config, vector_config = _runtime_sections(config)
    if not isinstance(query, str) or not query.strip():
        raise ValueError("A consulta não pode ser vazia.")
    if not isinstance(top_k, int) or isinstance(top_k, bool) or top_k <= 0:
        raise ValueError("top_k deve ser um inteiro positivo.")
    paths = _optional_mapping(config, "caminhos")
    return _retrieve(
        query.strip(),
        embedding_config,
        vector_config,
        project_id=project_id,
        tenant_id=tenant_id,
        k=top_k,
        persist_dir=paths.get("vector_store_dir"),
    )


_TOPK_CONHECIMENTO = 4
_TOPK_BASE = 6
_MENSAGEM_GENERICA = "Não foi possível obter a resposta do assistente no momento."


def _colecao_ausente(erro: Exception) -> bool:
    """Reconhece, sem importar SDK, a resposta 404 do vector store para coleção inexistente."""
    return getattr(erro, "status_code", None) == 404


def assistant_reply(
    *,
    llm_config: _Mapping[str, _Any],
    knowledge_config: _Mapping[str, _Any],
    historico: _Sequence[_Any],
    texto: str,
    base: _Mapping[str, _Any] | None = None,
    colecoes: _Any = None,
    estado: _Any = None,
    processado: _Any = None,
) -> _AssistantReply:
    """Responde um turno do assistente especialista em RAG (SPEC-0021).

    ``knowledge_config`` e ``base["config"]`` têm o formato de ``query_base``.
    ``base``, quando presente, é ``{"nome", "documentos", "config"}`` já validada
    no escopo da identidade. Qualquer falha do modelo, da recuperação ou do tempo
    limite sobe como ``AssistantError`` genérico, sem detalhe do provedor.
    """

    def buscar_conhecimento(pergunta: str):
        try:
            return query_base(knowledge_config, pergunta, top_k=_TOPK_CONHECIMENTO)
        except Exception as erro:
            if _colecao_ausente(erro):
                return None
            raise

    base_do_turno = None
    if base is not None:
        base_config = base["config"]
        base_do_turno = {
            "nome": base["nome"],
            "documentos": list(base.get("documentos") or []),
            "buscar": lambda pergunta: query_base(base_config, pergunta, top_k=_TOPK_BASE),
        }
    try:
        modelo = _ProviderFactory.create_chat_model(
            llm_config.get("provider", "openrouter"),
            {chave: valor for chave, valor in llm_config.items() if chave != "provider"},
        )
        return _executar_turno(
            chat_model=modelo,
            buscar_conhecimento=buscar_conhecimento,
            historico=historico,
            texto=texto,
            base=base_do_turno,
            colecoes=colecoes,
            estado=estado,
            processado=processado,
            timeout=llm_config.get("timeout"),
        )
    except Exception as erro:
        raise _AssistantError(_MENSAGEM_GENERICA) from erro


# Somente estes caminhos entram na coleção de conhecimento (SPEC-0021, DEC-005):
# nada de código-fonte, `.env`, specs ou dados de tenant.
_DOCUMENTOS_DO_ASSISTENTE = (
    "PROJECT.md",
    "INTERFACE.md",
    ".specsfy/STACK.md",
    ".specsfy/DATABASE.md",
)


def _documentos_do_assistente(raiz: _Path) -> list[_Path]:
    docs = sorted(caminho for caminho in (raiz / "docs").rglob("*.md") if caminho.is_file())
    fixos = [raiz / relativo for relativo in _DOCUMENTOS_DO_ASSISTENTE]
    return docs + [caminho for caminho in fixos if caminho.is_file()]


def index_assistant_knowledge(config: _Mapping[str, _Any], *, raiz: _Path | str) -> _SyncCounts:
    """Indexa a documentação viva do produto na coleção de conhecimento do assistente.

    Idempotente: ids de chunk são determinísticos e a sincronização em modo
    ``mirror`` mantém o que já existe e remove o que saiu da documentação. Cada
    chunk guarda o caminho de origem em ``caminho_origem`` para citação de fonte.
    """
    raiz = _Path(raiz)
    project_id, _embedding, vector_config = _runtime_sections(config)
    documentos: list[_Document] = []
    for caminho in _documentos_do_assistente(raiz):
        corpo = caminho.read_text(encoding="utf-8-sig").strip()
        if not corpo:
            continue
        relativo = caminho.relative_to(raiz).as_posix()
        documentos.append(
            _Document(
                page_content=corpo,
                metadata={
                    "nome_arquivo": relativo,
                    "caminho_relativo": relativo,
                    "caminho_origem": relativo,
                    "source_hash": _content_hash(corpo),
                },
            )
        )
    chunking = dict(_optional_mapping(config, "chunking"))
    chunks = _chunk_documents(documentos, chunking, project_id=project_id)
    destino = {**config, "vector_store": {"sync_mode": "mirror", **vector_config}}
    return sync_vector_store(config=destino, chunks=chunks)


def validar_config_conversao(config: _Mapping[str, _Any]) -> dict[str, _Any]:
    """Valida a configuração de conversão pela fachada neutra da aplicação."""
    from dataclasses import asdict

    return asdict(_carregar_config_conversao(config))


def convert_document(
    caminho: _Path,
    extensao: str,
    config: _Mapping[str, _Any] | None = None,
    abas: list[str] | None = None,
) -> _ConversionResult:
    """Converte um documento local para Markdown pela extensão informada (SPEC-0023)."""
    if abas is not None and extensao in {"xls", "xlsx"}:
        from src.pipeline.conversion.converters.xlsx import XlsxConverter
        from src.pipeline.conversion.legacy_office import normalizar_formato_legado

        if extensao == "xls":
            caminho = normalizar_formato_legado(caminho, "xls")
        return XlsxConverter(abas=abas).converter(caminho)
    conversor = _ConverterFactory.for_extension(
        extensao, config=_carregar_config_conversao(config or {})
    )
    return conversor.converter(caminho)


def validar_url_conversao(url: str) -> str:
    from src.pipeline.conversion.url_fetcher import url_sem_segredos, validar_url_publica

    validar_url_publica(url)
    return url_sem_segredos(url)


def buscar_url_conversao(url: str, config: _Mapping[str, _Any]):
    from src.pipeline.conversion.url_fetcher import UrlFetcher

    return UrlFetcher(_carregar_config_conversao(config)).buscar(url)


def listar_abas_conversao(caminho: _Path, extensao: str) -> list[str]:
    from openpyxl import load_workbook

    from src.pipeline.conversion.legacy_office import normalizar_formato_legado

    try:
        if extensao == "xls":
            caminho = normalizar_formato_legado(caminho, extensao)
        livro = load_workbook(caminho, read_only=True, data_only=True)
        try:
            if not livro.sheetnames:
                raise ValueError
            return list(livro.sheetnames)
        finally:
            livro.close()
    except Exception:  # noqa: BLE001 — falha pública segura de biblioteca/arquivo
        raise ConversionUserError(
            "Não foi possível ler as abas; confira a planilha e reenvie."
        ) from None


def avaliar_qualidade_conversao(markdown: str, config: _Mapping[str, _Any]):
    from src.pipeline.conversion.qualidade import avaliar_qualidade

    return avaliar_qualidade(markdown, _carregar_config_conversao(config))


def reprocessar_conversao_ia(
    caminho: _Path, formato: str, llm_config: _Mapping[str, _Any], config: _Mapping[str, _Any]
) -> str:
    from src.pipeline.conversion.reprocessar_ia import ReprocessadorIa

    class _ModeloSobPedido:
        def responder(self, mensagens, timeout=None):
            modelo = _ProviderFactory.create_chat_model(
                llm_config.get("provider", "openrouter"),
                {k: v for k, v in llm_config.items() if k != "provider"},
            )
            return modelo.responder(mensagens, timeout=timeout)

    return ReprocessadorIa(_ModeloSobPedido(), _carregar_config_conversao(config)).reprocessar(
        caminho, formato
    )


def _opcoes_limpeza(config: _Mapping[str, _Any] | None) -> dict[str, _Any]:
    return dict((config or {}).get("limpeza") or {})


def aplicar_regras_limpeza(markdown: str, arquivo: str):
    """Camada 1 da limpeza: regras fixas, sem IA (SPEC-0037)."""
    from src.pipeline.limpeza import regras

    return regras.aplicar_regras_fixas(markdown, arquivo)


def limpar_secoes_ia(
    markdown: str,
    arquivo: str,
    modelo: _Any,
    dominio: str,
    config: _Mapping[str, _Any] | None = None,
):
    """Camada 2: avaliação por IA com confiança; devolve (texto, exclusões, dúvidas)."""
    from src.pipeline.limpeza.avaliacao_ia import (
        LIMIAR_PADRAO,
        MAX_CHARS_PADRAO,
        AvaliadorIa,
        limpar_com_ia,
    )

    opcoes = _opcoes_limpeza(config)
    avaliador = AvaliadorIa(
        modelo,
        limiar=float(opcoes.get("limiar_confianca", LIMIAR_PADRAO)),
        max_chars_por_requisicao=int(opcoes.get("max_chars_requisicao", MAX_CHARS_PADRAO)),
    )
    return limpar_com_ia(markdown, arquivo, avaliador, dominio)


def deduplicar_documentos_limpeza(
    documentos: _Sequence[tuple[str, str]], config: _Mapping[str, _Any] | None = None
):
    """Camada de duplicidade entre os documentos do lote (SPEC-0037, FR-006)."""
    from src.pipeline.limpeza import duplicidade

    opcoes = _opcoes_limpeza(config)
    return duplicidade.deduplicar(
        documentos,
        minimo_caracteres=int(
            opcoes.get("minimo_caracteres_duplicidade", duplicidade.MINIMO_CARACTERES)
        ),
        limiar=float(opcoes.get("limiar_duplicidade", duplicidade.LIMIAR)),
        limiar_provavel=float(
            opcoes.get("limiar_duplicidade_provavel", duplicidade.LIMIAR_PROVAVEL)
        ),
    )


def propor_dominio_ia(amostra: _Sequence[str], modelo: _Any) -> str:
    from src.pipeline.limpeza.dominio import propor_dominio

    return propor_dominio(amostra, modelo)


def montar_relatorio_limpeza(
    documentos: _Sequence[_Mapping[str, _Any]],
    provaveis: _Sequence[_Any],
    *,
    ia_usada: bool,
    provedor: str | None,
    dominio_versao: int | None,
    motivo_sem_ia: str | None = None,
) -> str:
    from src.pipeline.limpeza.contratos import Exclusao
    from src.pipeline.limpeza.relatorio import DocumentoRelatorio, montar_relatorio

    itens = [
        DocumentoRelatorio(
            arquivo=d["arquivo"],
            exclusoes=[Exclusao(trecho="", **e) for e in d.get("exclusoes", [])],
            duvidas=list(d.get("duvidas", [])),
            revisar=list(d.get("revisar", [])),
            falha=d.get("falha"),
        )
        for d in documentos
    ]
    return montar_relatorio(
        itens,
        list(provaveis),
        ia_usada=ia_usada,
        provedor=provedor,
        dominio_versao=dominio_versao,
        motivo_sem_ia=motivo_sem_ia,
    )


def criar_modelo_chat(configuracao: _Mapping[str, _Any]):
    """Cria o adapter de chat da conexão de IA do Workspace; a chave só existe em memória."""
    dados = dict(configuracao)
    provider = dados.pop("provider", "openrouter")
    return _ProviderFactory.create_chat_model(provider, dados)


def extrair_afirmacoes(parte: _Sequence[tuple[str, str]], modelo: _Any) -> dict[str, _Any]:
    """Extrai afirmações de uma parte dos documentos (SPEC-0038, FR-002).

    Devolve ``{"afirmacoes": [dict], "nao_classificadas": int}``; a camada de aplicação
    só trabalha com dicionários simples.
    """
    from dataclasses import asdict

    from src.pipeline.contradicoes.afirmacoes import extrair_afirmacoes as _extrair

    resultado = _extrair(modelo, parte)
    return {
        "afirmacoes": [asdict(a) for a in resultado.afirmacoes],
        "nao_classificadas": resultado.nao_classificadas,
    }


def dividir_documentos_em_partes(documentos: _Mapping[str, str], tamanho: int):
    from src.pipeline.contradicoes.afirmacoes import dividir_em_partes

    return dividir_em_partes(documentos, tamanho)


def _afirmacoes(dados: _Sequence[_Mapping[str, _Any]]):
    from src.pipeline.contradicoes.contratos import Afirmacao

    return [
        Afirmacao(
            d["arquivo"], d["assunto"], d["texto"], d.get("impacto", "baixo"), d.get("imagem")
        )
        for d in dados
    ]


def agrupar_assuntos(
    afirmacoes: _Sequence[_Mapping[str, _Any]], maximo_por_julgamento: int
) -> list[dict[str, _Any]]:
    """Um julgamento por assunto normalizado, com somente as afirmações dele."""
    from dataclasses import asdict

    from src.pipeline.contradicoes.assuntos import planejar_julgamentos

    return [
        {"assunto": j.assunto, "afirmacoes": [asdict(a) for a in j.afirmacoes]}
        for j in planejar_julgamentos(_afirmacoes(afirmacoes), maximo_por_julgamento)
    ]


def julgar_assunto(
    modelo: _Any, assunto: str, afirmacoes: _Sequence[_Mapping[str, _Any]]
) -> list[dict[str, _Any]]:
    from dataclasses import asdict

    from src.pipeline.contradicoes.julgamento import julgar_assunto as _julgar

    return [asdict(o) for o in _julgar(modelo, assunto, _afirmacoes(afirmacoes))]


def montar_relatorio_contradicoes(
    analise: _Mapping[str, _Any], ocorrencias: _Sequence[_Mapping[str, _Any]], formato: str
):
    """Relatório em ``md`` (texto) ou ``pdf`` (bytes) com o mesmo conteúdo."""
    from types import SimpleNamespace

    from src.pipeline.contradicoes import relatorio

    itens = [
        SimpleNamespace(
            **{
                **o,
                "trechos": [SimpleNamespace(**{"imagem": None, **t}) for t in o["trechos"]],
                "arquivo_mais_antigo": o.get("arquivo_mais_antigo"),
            }
        )
        for o in ocorrencias
    ]
    resumo = SimpleNamespace(**analise)
    if formato == "pdf":
        return relatorio.montar_pdf(resumo, itens)
    return relatorio.montar_markdown(resumo, itens)


def interpretar_resposta(
    modelo: _Any,
    resposta: str,
    ocorrencias: _Sequence[_Mapping[str, _Any]],
    documentos: _Mapping[str, str],
) -> list[dict[str, _Any]]:
    """Propostas de correção (dicionários) a partir da resposta colada da empresa."""
    from dataclasses import asdict

    from src.pipeline.decisoes.interpretacao import interpretar_resposta as _interpretar

    return [asdict(p) for p in _interpretar(modelo, resposta, ocorrencias, documentos)]


def localizar_trecho(markdown: str, trecho: str) -> dict[str, _Any] | None:
    from dataclasses import asdict

    from src.pipeline.decisoes.localizacao import localizar_trecho as _localizar

    achado = _localizar(markdown, trecho)
    return asdict(achado) if achado else None


def aplicar_correcao(markdown: str, trecho_alvo: str, texto_corrigido: str) -> str:
    """Troca só o trecho alvo; ``ValueError`` quando ele não é localizado."""
    from src.pipeline.decisoes.aplicacao import aplicar_correcao as _aplicar

    return _aplicar(markdown, trecho_alvo, texto_corrigido)
