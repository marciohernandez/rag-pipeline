"""Conversor de XLSX para Markdown via openpyxl (SPEC-0023)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from src.pipeline.conversion.contracts import ConversionError, ConversionResult


def _celula_para_texto(valor: object) -> str:
    return "" if valor is None else str(valor)


def _planilha_para_tabela(linhas: list[tuple]) -> str | None:
    if not linhas:
        return None
    cabecalho, *dados = ([_celula_para_texto(v) for v in linha] for linha in linhas)
    partes = [
        f"| {' | '.join(cabecalho)} |",
        f"| {' | '.join('---' for _ in cabecalho)} |",
    ]
    partes.extend(f"| {' | '.join(linha)} |" for linha in dados)
    return "\n".join(partes)


class XlsxConverter:
    """Transforma cada planilha em uma seção Markdown com tabela."""

    def __init__(self, abas: list[str] | None = None):
        self.abas = abas

    def converter(self, caminho: Path) -> ConversionResult:
        try:
            pasta_de_trabalho = load_workbook(str(caminho), read_only=True, data_only=True)
        except Exception as erro:
            raise ConversionError(f"Falha ao converter XLSX: {caminho.name}") from erro

        try:
            if self.abas is not None and (
                not self.abas or not set(self.abas) <= set(pasta_de_trabalho.sheetnames)
            ):
                raise ConversionError("Selecione ao menos uma aba válida.")
            secoes: list[str] = []
            for planilha in pasta_de_trabalho.worksheets:
                if self.abas is not None and planilha.title not in self.abas:
                    continue
                cabecalho = f"## {planilha.title}"
                tabela = _planilha_para_tabela(list(planilha.iter_rows(values_only=True)))
                secoes.append("\n\n".join([cabecalho, tabela] if tabela else [cabecalho]))
            return ConversionResult(markdown="\n\n".join(secoes))
        finally:
            pasta_de_trabalho.close()
