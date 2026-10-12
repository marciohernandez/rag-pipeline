"""Conversor de PPTX para Markdown via python-pptx (SPEC-0023)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from src.pipeline.conversion.contracts import ConversionError, ConversionResult


class PptxConverter:
    """Transforma cada slide em uma seção Markdown com título e corpo."""

    def converter(self, caminho: Path) -> ConversionResult:
        try:
            apresentacao = Presentation(str(caminho))
        except Exception as erro:
            raise ConversionError(f"Falha ao converter PPTX: {caminho.name}") from erro

        secoes: list[str] = []
        for numero, slide in enumerate(apresentacao.slides, start=1):
            titulo_shape = slide.shapes.title
            tem_titulo = titulo_shape is not None and titulo_shape.has_text_frame
            titulo = titulo_shape.text.strip() if tem_titulo else ""
            cabecalho = f"## Slide {numero}: {titulo}" if titulo else f"## Slide {numero}"

            corpo = [
                shape.text_frame.text.strip()
                for shape in slide.shapes
                if shape is not titulo_shape
                and shape.has_text_frame
                and shape.text_frame.text.strip()
            ]
            secoes.append("\n\n".join([cabecalho, *corpo]))

        return ConversionResult(markdown="\n\n".join(secoes))
